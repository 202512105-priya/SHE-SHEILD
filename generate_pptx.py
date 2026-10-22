import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    C_PURPLE = RGBColor(88, 28, 135)     # Primary deep purple
    C_PURPLE_LIGHT = RGBColor(245, 240, 255) # Light purple card background
    C_MAGENTA = RGBColor(219, 39, 119)   # Accent magenta
    C_DARK = RGBColor(30, 41, 59)        # Slate dark text
    C_GRAY = RGBColor(100, 116, 139)     # Subtitle gray
    C_WHITE = RGBColor(255, 255, 255)
    C_BORDER = RGBColor(216, 180, 254)   # Soft lavender border
    C_TABLE_HEADER = RGBColor(76, 29, 149) # Table header dark purple
    C_TABLE_BG = RGBColor(250, 248, 255)

    def add_header(slide, title_text, category_text=None):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Arial"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = C_PURPLE
        if category_text:
            p2 = tf.add_paragraph()
            p2.text = category_text
            p2.font.name = "Arial"
            p2.font.size = Pt(14)
            p2.font.bold = False
            p2.font.color.rgb = C_MAGENTA

    def format_table(table, col_widths, font_size=Pt(11)):
        for i, w in enumerate(col_widths):
            table.columns[i].width = w
        for col_idx, cell in enumerate(table.rows[0].cells):
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_TABLE_HEADER
            for p in cell.text_frame.paragraphs:
                p.font.name = "Arial"
                p.font.size = Pt(12)
                p.font.bold = True
                p.font.color.rgb = C_WHITE
                p.alignment = PP_ALIGN.CENTER
        for row_idx in range(1, len(table.rows)):
            bg = C_TABLE_BG if row_idx % 2 == 1 else C_WHITE
            for col_idx, cell in enumerate(table.rows[row_idx].cells):
                cell.fill.solid()
                cell.fill.fore_color.rgb = bg
                for p in cell.text_frame.paragraphs:
                    p.font.name = "Arial"
                    p.font.size = font_size
                    p.font.color.rgb = C_DARK

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    
    # University Logo (Top Right)
    dau_logo = "extracted_assets/slide_1_shape_5.png"
    if os.path.exists(dau_logo):
        s1.shapes.add_picture(dau_logo, Inches(9.8), Inches(0.4), width=Inches(2.8))
    
    # Project Logo (Left)
    shield_logo = "extracted_assets/she_shield_logo.jpg"
    if os.path.exists(shield_logo):
        s1.shapes.add_picture(shield_logo, Inches(0.8), Inches(1.8), width=Inches(3.2))

    # Header course info
    tb_course = s1.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.5), Inches(1.1))
    tf_c = tb_course.text_frame
    tf_c.word_wrap = True
    p = tf_c.paragraphs[0]
    p.text = "WSSOA (IT644) Project Presentation I"
    p.font.name = "Arial"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = C_PURPLE
    p2 = tf_c.add_paragraph()
    p2.text = "Autumn 2026  ·  Dhirubhai Ambani University"
    p2.font.name = "Arial"
    p2.font.size = Pt(14)
    p2.font.color.rgb = C_GRAY

    # Project Title & Subtitle
    tb_title = s1.shapes.add_textbox(Inches(4.3), Inches(1.6), Inches(8.2), Inches(1.4))
    tf_t = tb_title.text_frame
    tf_t.word_wrap = True
    p = tf_t.paragraphs[0]
    p.text = "Project Title : SHE SHIELD"
    p.font.name = "Arial"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = C_PURPLE
    p_sub = tf_t.add_paragraph()
    p_sub.text = "Proactive AI-Powered Women Safety & Emergency Response Ecosystem"
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(15)
    p_sub.font.bold = True
    p_sub.font.color.rgb = C_MAGENTA
    p_desc = tf_t.add_paragraph()
    p_desc.text = "A cross-platform smart safety, real-time geo-tracking, and cryptographic evidence platform"
    p_desc.font.name = "Arial"
    p_desc.font.size = Pt(13)
    p_desc.font.color.rgb = C_GRAY

    # Team Members Table
    table_shape = s1.shapes.add_table(9, 3, Inches(4.3), Inches(3.2), Inches(8.2), Inches(3.8))
    t1 = table_shape.table
    members = [
        ("Sr. No.", "Student Name", "Student ID"),
        ("1", "Diya Jain", "202512014"),
        ("2", "Piyush Tanwani", "202512021"),
        ("3", "Divyesh Dandwani", "202512043"),
        ("4", "Harshil Dodwani", "202512044"),
        ("5", "Vivek Dhanwani", "202512045"),
        ("6", "Yash Gangwani", "202512048"),
        ("7", "Aman Rohera", "202512049"),
        ("8", "Priya Shah", "202512100")
    ]
    for r_idx, row in enumerate(members):
        for c_idx, val in enumerate(row):
            t1.cell(r_idx, c_idx).text = val
    format_table(t1, [Inches(1.2), Inches(4.5), Inches(2.5)], font_size=Pt(10.5))

    # ==========================================
    # SLIDE 2: PROBLEM DEFINITION
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "PROBLEM DEFINITION")

    # Column 1 Header: The Traditional Women's Safety Journey
    tb_c1 = s2.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(5.6), Inches(0.4))
    p = tb_c1.text_frame.paragraphs[0]
    p.text = "The Traditional Women's Safety Journey"
    p.font.name = "Arial"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = C_PURPLE

    journey_steps = [
        ("1. Commuting or Traveling Alone", "Navigating poorly lit roads, isolated transit routes, or suspicious cab rides late at night."),
        ("2. Encountering Distress or Threat", "Sensing imminent physical stalking, harassment, or unauthorized route diversions."),
        ("3. Attempting Emergency Actions", "Frantically unlocking a smartphone, dialing 112, or sending messages under intense panic."),
        ("4. Response Delays & Uncertainty", "Guardians and emergency services receive late alerts lacking precise GPS or live audio context."),
        ("5. Post-Incident Forensic Hurdles", "Evidence is erased, tampered with, or rejected in court due to broken chain-of-custody.")
    ]
    y_pos = 1.8
    for title, desc in journey_steps:
        box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y_pos), Inches(5.6), Inches(0.9))
        box.fill.solid()
        box.fill.fore_color.rgb = C_PURPLE_LIGHT
        box.line.color.rgb = C_BORDER
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.12)
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Arial"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = C_PURPLE
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = "Arial"
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_DARK
        y_pos += 1.02

    # Column 2 Header: Problems Faced During This Process
    tb_c2 = s2.shapes.add_textbox(Inches(6.8), Inches(1.3), Inches(5.7), Inches(0.4))
    p = tb_c2.text_frame.paragraphs[0]
    p.text = "Problems Faced During This Process"
    p.font.name = "Arial"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = C_MAGENTA

    problems = [
        ("1. High Panic Latency & Delayed Alerts", "Under physical duress, unlocking a phone or dialing emergency numbers is nearly impossible."),
        ("2. Route Deviation & Stalking Blindspots", "Traditional navigation apps do not monitor driver detours or notify guardians automatically."),
        ("3. Fragile & Tamperable Evidence", "Audio/video recordings stored on phones can be erased or challenged for lack of cryptographic proof."),
        ("4. Fragmented Responder Silos", "Guardians, nearby citizens, and civic ERSS-112 operate without a unified situational map."),
        ("5. Unchecked Cyber Stalking & Abuse", "Online harassment, blackmail, and extortion spread unchecked without rapid AI threat triage.")
    ]
    y_pos = 1.8
    for title, desc in problems:
        box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(y_pos), Inches(5.7), Inches(0.9))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(255, 245, 245)
        box.line.color.rgb = RGBColor(254, 202, 202)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.12)
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Arial"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(185, 28, 28)
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = "Arial"
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_DARK
        y_pos += 1.02

    # ==========================================
    # SLIDE 3: PROPOSED SOLUTION
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "PROPOSED SOLUTION")

    # Banner
    tb_b = s3.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(0.9))
    tf = tb_b.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "SheShield - Proactive AI-Powered Women Safety & Emergency Response Ecosystem"
    p.font.name = "Arial"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_PURPLE
    p2 = tf.add_paragraph()
    p2.text = "SheShield is an integrated, cross-platform safety ecosystem engineered to proactively prevent distress, dispatch multi-modal stealth SOS alerts in real time, monitor journeys with intelligent geofencing, and capture tamper-proof digital evidence."
    p2.font.name = "Arial"
    p2.font.size = Pt(13)
    p2.font.color.rgb = C_DARK

    # 4 Pillar Grid
    pillars = [
        ("1. Multi-Modal Emergency SOS & Stealth Triggers", 
         "• Trigger alerts via discreet BLE wearable button, secret voice phrase ('capcap'), or power/volume key double clicks.\n• Silent SOS mode dispatches alerts with zero screen illumination.\n• Offline coordinate caching with automated SMS fallback queue."),
        ("2. Real-Time Travel & Geo-Fencing Guardian",
         "• 'Cab Guardian' & 'Walk Guardian' monitor dynamic routes and GPS breadcrumbs.\n• Automated route deviation alerts when driver detours exceed threshold offsets.\n• AI Unsafe Zone Engine alerts users near high-risk nocturnal crime clusters."),
        ("3. Tamper-Proof Digital Evidence Vault",
         "• Automatic background audio/video recording upon travel start or SOS trigger.\n• Instant cryptographic SHA-256 digital fingerprinting of evidence assets.\n• Tamper-verification pass ensuring forensic chain-of-custody admissible in court."),
        ("4. Unified Multi-Role Command Network",
         "• Dedicated real-time dashboards for At-Risk Users, Linked Family Guardians, ERSS-112 Emergency Responders, and Police Administrators.\n• Live WebSocket coordinate streaming and automated incident dossier export.\n• Cyber Crime reporting suite with automated threat risk assessment.")
    ]
    positions = [
        (Inches(0.8), Inches(2.2), Inches(5.7), Inches(2.2)),
        (Inches(6.8), Inches(2.2), Inches(5.7), Inches(2.2)),
        (Inches(0.8), Inches(4.55), Inches(5.7), Inches(2.1)),
        (Inches(6.8), Inches(4.55), Inches(5.7), Inches(2.1))
    ]
    for i, (title, content) in enumerate(pillars):
        x, y, w, h = positions[i]
        box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        box.fill.solid()
        box.fill.fore_color.rgb = C_PURPLE_LIGHT
        box.line.color.rgb = C_BORDER
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.15)
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Arial"
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = C_PURPLE
        p2 = tf.add_paragraph()
        p2.text = content
        p2.font.name = "Arial"
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = C_DARK

    # Bottom Expected Outcome Card
    out_box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.75), Inches(11.7), Inches(0.6))
    out_box.fill.solid()
    out_box.fill.fore_color.rgb = RGBColor(238, 242, 255)
    out_box.line.color.rgb = RGBColor(199, 210, 254)
    tf = out_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Expected Outcome : SheShield compresses emergency response latency from minutes to seconds, prevents route deviations proactively, empowers women with stealth triggers, and provides incontrovertible legal evidence for swift justice."
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(49, 46, 129)

    # ==========================================
    # SLIDE 4: SYSTEM WORKFLOW
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "SYSTEM WORKFLOW")

    # Image Diagram (Left)
    wf_img = "extracted_assets/she_shield_workflow.png"
    if os.path.exists(wf_img):
        s4.shapes.add_picture(wf_img, Inches(0.8), Inches(1.3), width=Inches(5.0))

    # Text explanation (Right)
    tb_wf = s4.shapes.add_textbox(Inches(6.1), Inches(1.3), Inches(6.4), Inches(5.8))
    tf = tb_wf.text_frame
    tf.word_wrap = True
    wf_points = [
        ("Data Integration & Telemetry Ingestion:", "Collects multi-modal inputs from BLE wearable panic buttons, secret voice phrase detection, volume/power hardware keys, and GPS sensors into a unified telemetry stream."),
        ("Autonomous Edge Telemetry & Offline Caching:", "Caches high-frequency GPS breadcrumbs locally during cellular drops and queues automated SMS fallback alerts with Google Maps coordinates."),
        ("Geo-Spatial Anomaly & Deviation Engine:", "Continuously evaluates live transit coordinates against planned route vectors; flags unauthorized route offsets exceeding threshold limits (e.g. 30m+)."),
        ("AI Unsafe Zone Intelligence:", "Assesses real-time user location against nocturnal risk factors, street illumination, and historical crime complaints to recommend safe corridors."),
        ("Cryptographic Forensic Preservation:", "Streams background audio/video into the Evidence Vault, calculating SHA-256 cryptographic digests immediately to guarantee court-admissible chain of custody."),
        ("Unified Command & Multi-Role Dispatch:", "Dispatches real-time incident notifications to linked family guardians and provides ERSS-112 police dispatchers with live tracking and incident dossiers.")
    ]
    for i, (title, desc) in enumerate(wf_points):
        p1 = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Arial"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = C_PURPLE
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = "Arial"
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_DARK
        p_space = tf.add_paragraph()
        p_space.text = ""
        p_space.font.size = Pt(3)

    # ==========================================
    # SLIDE 5: PROJECT SCOPE
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "PROJECT SCOPE")

    # In Scope Box
    box_in = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(5.6), Inches(4.7))
    box_in.fill.solid()
    box_in.fill.fore_color.rgb = C_PURPLE_LIGHT
    box_in.line.color.rgb = C_BORDER
    tf = box_in.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "In Scope :"
    p.font.name = "Arial"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = C_PURPLE

    in_items = [
        "Multi-modal emergency SOS triggering (BLE wearable button, secret voice phrase 'capcap', power/volume double click, UI button).",
        "Real-time GPS tracking, breadcrumb history, and offline location caching with automated SMS fallback queue.",
        "Cab Guardian & Walk Guardian modules with automated route deviation alerts and distance-offset detection.",
        "AI Unsafe Zone Intelligence calculating localized risk scores based on time, lighting, and historical crime clusters.",
        "Cryptographic Evidence Locker with SHA-256 digital fingerprinting for tamper-proof audio/video blackbox files.",
        "Role-based portals: User Dashboard, Family Guardian Portal, Emergency Responder Hub, and Police Admin Command Center.",
        "Cyber Crime Reporting suite with AI threat severity scoring, screenshot upload, and official complaint tracking.",
        "Architecture-level integration hooks for ERSS-112, CCTNS police records, and CERT-In cyber threat databases."
    ]
    for item in in_items:
        p_item = tf.add_paragraph()
        p_item.text = "• " + item
        p_item.font.name = "Arial"
        p_item.font.size = Pt(11)
        p_item.font.color.rgb = C_DARK

    # Out of Scope Box
    box_out = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.3), Inches(5.7), Inches(4.7))
    box_out.fill.solid()
    box_out.fill.fore_color.rgb = RGBColor(254, 242, 242)
    box_out.line.color.rgb = RGBColor(254, 202, 202)
    tf = box_out.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "Out of Scope :"
    p.font.name = "Arial"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = RGBColor(185, 28, 28)

    out_items = [
        "Physical deployment or tactical command of on-ground armed police personnel.",
        "Proprietary semiconductor or custom IoT hardware chip fabrication (utilizes standard Bluetooth LE and Android/iOS standards).",
        "Direct courtroom litigation, judicial prosecution, or formal legal sentencing.",
        "Guaranteed 100% real-time tracking in deep subterranean basements without GNSS, cellular GSM/SMS, or BLE relay connectivity.",
        "Automated facial recognition matching on low-resolution or heavily occluded third-party CCTV streams."
    ]
    for item in out_items:
        p_item = tf.add_paragraph()
        p_item.text = "• " + item
        p_item.font.name = "Arial"
        p_item.font.size = Pt(12)
        p_item.font.color.rgb = C_DARK

    # Project Boundary Box
    b_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.9))
    b_box.fill.solid()
    b_box.fill.fore_color.rgb = RGBColor(241, 245, 249)
    b_box.line.color.rgb = RGBColor(203, 213, 225)
    tf = b_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Project Boundary :  SheShield focuses on rapid threat detection, automated alerting, geo-spatial intelligence, and tamper-proof evidence preservation; physical rescue and law enforcement remain under authorized civic agencies."
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(30, 41, 59)

    # ==========================================
    # SLIDE 6: SYSTEM ARCHITECTURE
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "SYSTEM ARCHITECTURE")

    # Architecture Image (Top / Center)
    arch_img = "extracted_assets/she_shield_architecture.png"
    if os.path.exists(arch_img):
        s6.shapes.add_picture(arch_img, Inches(0.8), Inches(1.3), width=Inches(6.8))

    # Architecture Layers Description (Right)
    tb_arch = s6.shapes.add_textbox(Inches(7.8), Inches(1.3), Inches(4.7), Inches(5.8))
    tf = tb_arch.text_frame
    tf.word_wrap = True
    arch_details = [
        ("Presentation Layer:", "Flutter 3.x with Dart delivers 60fps responsive interfaces across mobile and web with Material 3 theming and modern glassmorphism."),
        ("Backend API Layer:", "Python Flask microservices orchestrate RESTful endpoints, token authentication, and Firebase Admin SDK services."),
        ("Application Services:", "Core micro-engines power Cab Deviation Detection, AI Unsafe Zone Scoring, Forensic Evidence Streaming, and SMS fallback queues."),
        ("Data Layer:", "Cloud Firestore provides real-time reactive NoSQL collections for live telemetry, paired with encrypted Firebase Storage."),
        ("Integration Layer:", "Future-ready integration hooks interface with ERSS-112 (National Emergency), CCTNS (Police records), CERT-In (Cyber Defense), and Twilio SMS.")
    ]
    for i, (title, desc) in enumerate(arch_details):
        p1 = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Arial"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = C_PURPLE
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = "Arial"
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_DARK
        p_space = tf.add_paragraph()
        p_space.text = ""
        p_space.font.size = Pt(3)

    # ==========================================
    # SLIDE 7: FRONTEND AND BACKEND TECHNOLOGIES
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "FRONTEND AND BACKEND TECHNOLOGIES")

    # Left: Frontend Table
    tb_fe_title = s7.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(5.6), Inches(0.4))
    p = tb_fe_title.text_frame.paragraphs[0]
    p.text = "FRONTEND TECHNOLOGIES"
    p.font.name = "Arial"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_PURPLE

    fe_table_shape = s7.shapes.add_table(7, 2, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2))
    fe_t = fe_table_shape.table
    fe_data = [
        ("Technology", "Why It Is Used in SheShield"),
        ("Flutter", "Builds high-performance native user interfaces for Android and iOS using a single unified codebase."),
        ("Dart", "Provides sound type safety and asynchronous event streams essential for emergency telemetry."),
        ("Provider", "Manages reactive state cleanly across SOS countdowns, travel monitoring, and live location updates."),
        ("Material 3", "Creates accessible, high-contrast UI components with responsive glassmorphic cards and buttons."),
        ("Flutter BLE", "Enables low-latency communication with Bluetooth LE panic buttons and wearable smart pendants."),
        ("Geolocator", "Captures high-precision GPS coordinates, breadcrumb paths, and speed telemetry in real time.")
    ]
    for r_idx, row in enumerate(fe_data):
        for c_idx, val in enumerate(row):
            fe_t.cell(r_idx, c_idx).text = val
    format_table(fe_t, [Inches(1.6), Inches(4.0)], font_size=Pt(10.5))

    # Right: Backend Table
    tb_be_title = s7.shapes.add_textbox(Inches(6.8), Inches(1.2), Inches(5.7), Inches(0.4))
    p = tb_be_title.text_frame.paragraphs[0]
    p.text = "BACKEND & AI TECHNOLOGIES"
    p.font.name = "Arial"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_MAGENTA

    be_table_shape = s7.shapes.add_table(7, 2, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2))
    be_t = be_table_shape.table
    be_data = [
        ("Technology", "Why It Is Used in SheShield"),
        ("Python Flask", "Runs lightweight, scalable REST microservices handling emergency routing and alert dispatch."),
        ("Flask-CORS", "Secures cross-origin API communication between mobile clients, responder maps, and admin consoles."),
        ("Firebase Admin SDK", "Provides privileged server-side access to Cloud Firestore collections and user security rules."),
        ("PyTorch & Transformers", "Powers NLP models to classify cyber harassment severity and assess digital blackmail threats."),
        ("OpenCV", "Processes uploaded forensic evidence screenshots and inspects visual file authenticity."),
        ("hashlib (SHA-256)", "Computes immutable cryptographic digests of recorded evidence to guarantee court admissibility.")
    ]
    for r_idx, row in enumerate(be_data):
        for c_idx, val in enumerate(row):
            be_t.cell(r_idx, c_idx).text = val
    format_table(be_t, [Inches(1.8), Inches(3.9)], font_size=Pt(10.5))

    # ==========================================
    # SLIDE 8: DATABASE, AI, DATA PROCESSING & INTEGRATION
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "DATABASE, AI, DATA PROCESSING & INTEGRATION")

    db_table_shape = s8.shapes.add_table(13, 3, Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.7))
    dbt = db_table_shape.table
    db_data = [
        ("Technology / Component", "Category", "Why It Is Used in SheShield"),
        ("Cloud Firestore", "Data Storage & Real-Time Sync", "Stores user profiles, SOS alerts, travel sessions, and syncs live telemetry across all roles."),
        ("Firebase Auth", "Authentication & Identity", "Secures multi-role authentication (User, Guardian, Responder, Admin) with biometric support."),
        ("Python (Flask)", "Backend Microservices", "Coordinates emergency dispatch, evidence hashing verification, and external API gateways."),
        ("PyTorch & Transformers", "AI Threat Intelligence", "Evaluates cyber crime complaints using NLP to calculate threat scores (e.g. 92% blackmail)."),
        ("AI Unsafe Zone Service", "Geo-Spatial AI Analytics", "Computes dynamic composite risk scores using time of night, illumination, and crime records."),
        ("Cab Deviation Engine", "Heuristic Route Analytics", "Calculates vehicle route offsets from expected polyline and triggers deviation alerts."),
        ("SHA-256 Crypto Engine", "Evidence Integrity & Forensics", "Generates cryptographic fingerprints for captured audio/video to ensure court admissibility."),
        ("ERSS-112 Integration Hook", "National Emergency Integration", "Connects SheShield alert payloads with India's Emergency Response Support System (112)."),
        ("CCTNS Sync Hook", "Law Enforcement Integration", "Prepares crime dossiers and incident records for Crime and Criminal Tracking Network & Systems."),
        ("CERT-In Webhook Hook", "Cyber Threat Intelligence", "Automates escalation of severe online blackmail, doxxing, and cyber harassment cases."),
        ("Twilio SMS Fallback", "Redundant Emergency Alerting", "Dispatches instant SMS with live Google Maps coordinates when mobile data is offline."),
        ("Docker Compose", "Service Management", "Packages and runs application microservices, databases, and mock government gateways.")
    ]
    for r_idx, row in enumerate(db_data):
        for c_idx, val in enumerate(row):
            dbt.cell(r_idx, c_idx).text = val
    format_table(dbt, [Inches(2.5), Inches(2.8), Inches(6.4)], font_size=Pt(9.5))

    # ==========================================
    # SLIDE 9: DATABASE DESIGN AND DATA MANAGEMENT
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "DATABASE DESIGN AND DATA MANAGEMENT")

    # Image Diagram (Left)
    db_img = "extracted_assets/she_shield_database_design.png"
    if os.path.exists(db_img):
        s9.shapes.add_picture(db_img, Inches(0.8), Inches(1.3), width=Inches(5.0))

    # Table of Collections (Right)
    coll_table_shape = s9.shapes.add_table(9, 2, Inches(6.0), Inches(1.3), Inches(6.5), Inches(5.7))
    ct = coll_table_shape.table
    coll_data = [
        ("Table / Collection Name", "Purpose in SheShield"),
        ("User", "Stores user account details, Aadhaar/face verification status, and safety preferences."),
        ("Guardians", "Maintains linked emergency contacts, kinship relations, and notification priority order."),
        ("SOSEvents", "Logs emergency alerts, trigger mechanisms (voice/BLE/button), live GPS, and status."),
        ("TravelTrips", "Records Cab Guardian & Walk Guardian trips, vehicle plates, route offsets, and risk scores."),
        ("EvidenceVault", "Catalogs encrypted audio/video assets alongside SHA-256 hashes for legal chain of custody."),
        ("CyberComplaints", "Archives cyber abuse reports, screenshots, AI threat severity scores, and police status."),
        ("UnsafeZones", "Defines geofenced high-risk corridors, historical incident density, and advisory rules."),
        ("AuditLogs", "Preserves an immutable chronological ledger of administrative actions and SOS events.")
    ]
    for r_idx, row in enumerate(coll_data):
        for c_idx, val in enumerate(row):
            ct.cell(r_idx, c_idx).text = val
    format_table(ct, [Inches(2.0), Inches(4.5)], font_size=Pt(10.5))

    # ==========================================
    # SLIDE 10: SMART SAFETY INTELLIGENCE AND AI CAPABILITIES
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "SMART SAFETY INTELLIGENCE AND AI CAPABILITIES")

    # Flowchart Diagram (Top / Center)
    intel_img = "extracted_assets/she_shield_intelligence.png"
    if os.path.exists(intel_img):
        s10.shapes.add_picture(intel_img, Inches(0.8), Inches(1.3), width=Inches(11.7))

    # Summary Paragraph (Bottom)
    box_summary = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.3), Inches(11.7), Inches(1.7))
    box_summary.fill.solid()
    box_summary.fill.fore_color.rgb = C_PURPLE_LIGHT
    box_summary.line.color.rgb = C_BORDER
    tf = box_summary.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "SheShield unifies multi-modal sensor intelligence, predictive geo-spatial analytics, NLP threat scoring, and cryptographic evidence preservation to build a proactive safety net. It eliminates panic latency through stealth triggers, proactively detects travel anomalies before danger strikes, and guarantees uncompromised digital evidence for swift law enforcement and court action."
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.color.rgb = C_DARK
    p_sub = tf.add_paragraph()
    p_sub.text = "Key Intelligence Pillars: Multi-Modal Stealth Triggers (BLE & Voice 'capcap') · Cab Deviation Analytics · AI Unsafe Zone Scoring · Forensic SHA-256 Evidence Vault · Automated Cyber Threat Assessment."
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(12)
    p_sub.font.bold = True
    p_sub.font.color.rgb = C_PURPLE

    # ==========================================
    # SLIDE 11: CHALLENGES AND LIMITATIONS
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "CHALLENGES AND LIMITATIONS")

    # Left: Challenges Faced Table
    tb_ch_title = s11.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(6.0), Inches(0.4))
    p = tb_ch_title.text_frame.paragraphs[0]
    p.text = "Challenges Faced :"
    p.font.name = "Arial"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_PURPLE

    ch_table_shape = s11.shapes.add_table(7, 2, Inches(0.8), Inches(1.7), Inches(6.0), Inches(5.2))
    cht = ch_table_shape.table
    ch_data = [
        ("Challenge", "Description"),
        ("Low / Intermittent Connectivity", "Emergency situations often occur in low-signal areas; mitigated via offline coordinate caching and automated SMS fallback queue."),
        ("Accidental / False Triggers", "Pocket presses could overburden responders; handled using 5-second cancellable audio-haptic countdown timers."),
        ("Urban GPS Signal Drift", "Multipath reflections cause coordinate jumping; stabilized using Kalman filtering and route offset tolerance thresholds."),
        ("Mobile OS Background Throttling", "Operating systems suspend background apps; resolved via Android foreground services and travel wake-locks."),
        ("Evidence Cryptographic Integrity", "Audio/video evidence must be court-admissible; solved by generating client-side SHA-256 hashes prior to cloud sync."),
        ("Multi-Stakeholder Dispatch Latency", "Synchronizing guardians, nearby responders, and police without confusion requires strict role-based state machines.")
    ]
    for r_idx, row in enumerate(ch_data):
        for c_idx, val in enumerate(row):
            cht.cell(r_idx, c_idx).text = val
    format_table(cht, [Inches(2.0), Inches(4.0)], font_size=Pt(9.5))

    # Right: System Limitations
    tb_lim_title = s11.shapes.add_textbox(Inches(7.2), Inches(1.2), Inches(5.3), Inches(0.4))
    p = tb_lim_title.text_frame.paragraphs[0]
    p.text = "System Limitations :"
    p.font.name = "Arial"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_MAGENTA

    box_lim = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.2), Inches(1.7), Inches(5.3), Inches(5.2))
    box_lim.fill.solid()
    box_lim.fill.fore_color.rgb = RGBColor(254, 242, 242)
    box_lim.line.color.rgb = RGBColor(254, 202, 202)
    tf = box_lim.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.2)
    
    limits = [
        ("Hardware Dependency:", "Stealth BLE wearable triggers depend on Bluetooth peripheral battery level and active pairing range."),
        ("Cellular SMS Network Latency:", "Delivery speed of emergency SMS alerts depends on telecom carrier traffic and congestion."),
        ("Govt Emergency Gateway Readiness:", "Direct dispatch into civic ERSS-112 and CCTNS relies on government API availability and authorization."),
        ("Ambient Acoustic Noise:", "Extreme environmental noise (e.g. industrial areas) can occasionally impede whisper-level voice recognition."),
        ("Polyline Granularity:", "Route deviation alerts require baseline map routing data and destination accuracy.")
    ]
    for i, (title, desc) in enumerate(limits):
        p1 = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p1.text = "• " + title
        p1.font.name = "Arial"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(185, 28, 28)
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = "Arial"
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_DARK
        p_sp = tf.add_paragraph()
        p_sp.text = ""
        p_sp.font.size = Pt(3)

    # ==========================================
    # SLIDE 12: THANK YOU
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    
    # Dhirubhai Ambani University Logo
    if os.path.exists(dau_logo):
        s12.shapes.add_picture(dau_logo, Inches(4.7), Inches(1.2), width=Inches(3.8))
    
    # Thank You text
    tb_ty = s12.shapes.add_textbox(Inches(2.0), Inches(2.8), Inches(9.3), Inches(1.2))
    tf = tb_ty.text_frame
    p = tf.paragraphs[0]
    p.text = "THANK YOU"
    p.font.name = "Arial"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = C_PURPLE
    p.alignment = PP_ALIGN.CENTER

    tb_sub = s12.shapes.add_textbox(Inches(2.0), Inches(4.0), Inches(9.3), Inches(2.2))
    tf = tb_sub.text_frame
    p = tf.paragraphs[0]
    p.text = "SHE SHIELD — Proactive AI-Powered Women Safety & Emergency Response Ecosystem"
    p.font.name = "Arial"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_MAGENTA
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = "WSSOA (IT644)  ·  Autumn 2026\nDhirubhai Ambani University, Gandhinagar, Gujarat\n\nQuestions & Answers"
    p2.font.name = "Arial"
    p2.font.size = Pt(14)
    p2.font.color.rgb = C_GRAY
    p2.alignment = PP_ALIGN.CENTER

    output_path = "SheShield_Project_Presentation.pptx"
    prs.save(output_path)
    print(f"Successfully generated {output_path} with {len(prs.slides)} slides!")

if __name__ == "__main__":
    build_presentation()
