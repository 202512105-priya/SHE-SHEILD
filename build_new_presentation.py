import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Elegant Color Palette
    C_PURPLE = RGBColor(76, 29, 149)      # Primary deep purple (#4C1D95)
    C_PURPLE_LIGHT = RGBColor(245, 240, 255) # Light purple card background
    C_MAGENTA = RGBColor(190, 24, 93)     # Accent magenta (#BE185D)
    C_DARK = RGBColor(30, 41, 59)         # Slate text
    C_GRAY = RGBColor(100, 116, 139)      # Subtitle gray
    C_WHITE = RGBColor(255, 255, 255)
    C_BORDER = RGBColor(216, 180, 254)    # Lavender border
    C_TABLE_HEADER = RGBColor(88, 28, 135)

    def add_header(slide, title_text, category_text=None):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Arial"
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = C_PURPLE
        if category_text:
            p2 = tf.add_paragraph()
            p2.text = category_text
            p2.font.name = "Arial"
            p2.font.size = Pt(13)
            p2.font.bold = False
            p2.font.color.rgb = C_MAGENTA

    def format_table(table, col_widths, font_size=Pt(10), header_bg=C_TABLE_HEADER):
        for i, w in enumerate(col_widths):
            table.columns[i].width = w
        for cell in table.rows[0].cells:
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_bg
            for p in cell.text_frame.paragraphs:
                p.font.name = "Arial"
                p.font.size = Pt(11)
                p.font.bold = True
                p.font.color.rgb = C_WHITE
                p.alignment = PP_ALIGN.CENTER
        for row_idx in range(1, len(table.rows)):
            bg = C_PURPLE_LIGHT if row_idx % 2 == 1 else C_WHITE
            for cell in table.rows[row_idx].cells:
                cell.fill.solid()
                cell.fill.fore_color.rgb = bg
                for p in cell.text_frame.paragraphs:
                    p.font.name = "Arial"
                    p.font.size = font_size
                    p.font.color.rgb = C_DARK

    # ==========================================
    # SLIDE 1: Title & Team Credentials
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    dau_logo = "extracted_assets/slide_1_shape_5.png"
    if os.path.exists(dau_logo):
        s1.shapes.add_picture(dau_logo, Inches(9.8), Inches(0.4), width=Inches(2.8))

    shield_logo = "extracted_assets/she_shield_logo.jpg"
    if os.path.exists(shield_logo):
        s1.shapes.add_picture(shield_logo, Inches(0.8), Inches(1.8), width=Inches(3.2))

    tb_course = s1.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.5), Inches(1.1))
    tf_c = tb_course.text_frame
    tf_c.word_wrap = True
    p = tf_c.paragraphs[0]
    p.text = "WSSOA (IT644) Project Presentation I  ·  Autumn 2026"
    p.font.name = "Arial"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = C_PURPLE
    p2 = tf_c.add_paragraph()
    p2.text = "Dhirubhai Ambani University, Gandhinagar"
    p2.font.name = "Arial"
    p2.font.size = Pt(14)
    p2.font.color.rgb = C_GRAY

    tb_title = s1.shapes.add_textbox(Inches(4.3), Inches(1.5), Inches(8.2), Inches(1.5))
    tf_t = tb_title.text_frame
    tf_t.word_wrap = True
    p = tf_t.paragraphs[0]
    p.text = "SHE SHIELD"
    p.font.name = "Arial"
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = C_PURPLE
    p_sub = tf_t.add_paragraph()
    p_sub.text = "Cyber-Integrated Safety Platform for Women"
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(16)
    p_sub.font.bold = True
    p_sub.font.color.rgb = C_MAGENTA
    p_desc = tf_t.add_paragraph()
    p_desc.text = "A Proactive, Multi-Modal Emergency Dispatch and Cryptographic Forensic Ecosystem"
    p_desc.font.name = "Arial"
    p_desc.font.size = Pt(12)
    p_desc.font.color.rgb = C_GRAY

    # Team Members & Roles Table
    table_shape = s1.shapes.add_table(9, 3, Inches(4.3), Inches(3.2), Inches(8.2), Inches(3.8))
    t1 = table_shape.table
    team_data = [
        ("Sr. No.", "Team Member", "Specific Engineering Role"),
        ("1", "Krisha Doshi", "Mobile UI Engineering (Flutter / Material 3)"),
        ("2", "Priya Shah", "Navigation & User Flow (Routing & State)"),
        ("3", "Jainam Shah", "Backend APIs (Flask REST Microservices)"),
        ("4", "Maahir Shah", "Database (PostgreSQL / Redis / Firestore)"),
        ("5", "Hitarth Shah", "AI Threat Models (PyTorch & NLP Classifier)"),
        ("6", "Jay Shah", "DevOps & Integration (Docker, CI/CD & Gateways)"),
        ("7", "Nirmal Patel", "Auth / RBAC (JWT & Role Permissions)"),
        ("8", "Jenil Shah", "Cybersecurity (SHA-256 Vault & AES Encryption)")
    ]
    for r_idx, row in enumerate(team_data):
        for c_idx, val in enumerate(row):
            t1.cell(r_idx, c_idx).text = val
    format_table(t1, [Inches(1.0), Inches(3.2), Inches(4.0)], font_size=Pt(10.5))

    # ==========================================
    # SLIDE 2: Problem Definition & Market Gaps
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "PROBLEM DEFINITION & MARKET GAPS", "Critical Limitations in Existing Personal Safety Solutions")

    prob_cards = [
        ("1. Single-Vector Focus (Physical vs Cyber)",
         "• Current safety apps only handle physical panic buttons and completely ignore cyber harassment.\n• Digital stalking, image morphing, doxxing, sextortion, and extortion affect millions of women daily.\n• Zero integration exists between physical transit monitors and cyber incident reporting desks.",
         RGBColor(254, 242, 242), RGBColor(248, 113, 113), RGBColor(185, 28, 28)),
        ("2. Acute Stress & Panic Latency",
         "• Victims under acute physical or psychological stress cannot unlock phones or navigate complex apps.\n• Requiring a user to illuminate a screen, input PIN/biometrics, and dial numbers alerts assailants.\n• Lack of stealth triggers (BLE pendants, secret voice phrases, tactile hardware buttons) causes fatal delays.",
         RGBColor(255, 247, 237), RGBColor(251, 146, 60), RGBColor(194, 65, 12)),
        ("3. Digital Evidence Destruction",
         "• Digital evidence (threats, blackmail, stalking chats) is easily deleted before police reporting.\n• Attackers force victims to wipe devices or utilize disappearing message features on social platforms.\n• Smartphone screenshots lack cryptographic timestamps, metadata, and hash proofs for Section 65B compliance.",
         RGBColor(253, 242, 248), RGBColor(244, 114, 182), RGBColor(190, 24, 93))
    ]
    y_pos = 1.3
    for title, content, bg, border, text_c in prob_cards:
        box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y_pos), Inches(11.7), Inches(1.55))
        box.fill.solid()
        box.fill.fore_color.rgb = bg
        box.line.color.rgb = border
        box.line.width = Pt(1.5)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.12)
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Arial"
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = text_c
        p2 = tf.add_paragraph()
        p2.text = content
        p2.font.name = "Arial"
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_DARK
        y_pos += 1.68

    # Bottom Callout Box
    out_b = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.45), Inches(11.7), Inches(0.65))
    out_b.fill.solid()
    out_b.fill.fore_color.rgb = RGBColor(238, 242, 255)
    out_b.line.color.rgb = RGBColor(199, 210, 254)
    tf = out_b.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Market Gap Summary : Modern safety necessitates a unified platform bridging immediate physical stealth triggers, proactive transit geofencing, and forensically sound cyber evidence preservation."
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(49, 46, 129)

    # ==========================================
    # SLIDE 3: Proposed Solution — The SHE SHIELD Ecosystem
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "PROPOSED SOLUTION — THE SHE SHIELD ECOSYSTEM", "Tri-Pillar Cyber-Physical Protection Architecture")

    sol_cards = [
        ("Pillar 1: Physical Emergency SOS + Real-Time Transit Tracking",
         "• Multi-modal stealth triggers: Bluetooth LE wearable panic pendant, secret phrase ('capcap'), power/volume click.\n• Sub-2-second emergency dispatch broadcasting GPS breadcrumbs to guardians & police consoles simultaneously.\n• 'Cab Guardian' & 'Walk Guardian' modules actively monitoring route deviations (>30m offset) and nocturnal hotspots.",
         Inches(0.8), Inches(1.3), Inches(11.7), Inches(1.55), (245, 243, 255), (167, 139, 250), (76, 29, 149)),
        ("Pillar 2: AI-Powered Cybercrime Desk with OCR Screenshot Scanning",
         "• Built-in OCR ingestion pipeline parsing screenshots of blackmail, extortion, doxxing, and cyber harassment chats.\n• PyTorch NLP threat model evaluating harassment severity and classifying incident danger (e.g. 92% Critical Risk).\n• Automated generation of standardized police complaint dossiers with extracted timestamps, sender handles, and evidence URLs.",
         Inches(0.8), Inches(3.0), Inches(11.7), Inches(1.55), (253, 244, 255), (232, 121, 249), (112, 26, 117)),
        ("Pillar 3: Tamper-Proof Cryptographic Evidence Vault",
         "• Continuous background blackbox audio/video recording initiated upon travel or emergency trigger.\n• Client-side SHA-256 cryptographic hashing executed prior to cloud upload, creating an immutable custody ledger.\n• AES-256-GCM encryption at rest and in transit, establishing legally incontrovertible proof for judicial court admissibility.",
         Inches(0.8), Inches(4.7), Inches(11.7), Inches(1.55), (238, 242, 255), (129, 140, 248), (49, 46, 129))
    ]
    for title, content, x, y, w, h, bg, border, text_c in sol_cards:
        box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(*bg)
        box.line.color.rgb = RGBColor(*border)
        box.line.width = Pt(1.5)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.12)
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Arial"
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(*text_c)
        p2 = tf.add_paragraph()
        p2.text = content
        p2.font.name = "Arial"
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_DARK

    out_s3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.45), Inches(11.7), Inches(0.65))
    out_s3.fill.solid()
    out_s3.fill.fore_color.rgb = RGBColor(240, 253, 250)
    out_s3.line.color.rgb = RGBColor(45, 212, 191)
    tf = out_s3.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Ecosystem Outcome : Unifies personal physical transit security with digital cyber defense into a single, high-availability mobile and cloud architecture."
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 118, 110)

    # ==========================================
    # SLIDE 4: Stakeholders & User Roles (RBAC)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "STAKEHOLDERS & USER ROLES (RBAC)", "Role-Based Access Control and Persona Matrix")

    rbac_img = "extracted_assets/she_shield_stakeholders_rbac.png"
    if os.path.exists(rbac_img):
        s4.shapes.add_picture(rbac_img, Inches(0.8), Inches(1.3), width=Inches(11.7))

    # Summary table below image
    rb_table_shape = s4.shapes.add_table(5, 3, Inches(0.8), Inches(4.5), Inches(11.7), Inches(2.5))
    rbt = rb_table_shape.table
    rb_data = [
        ("Role Persona", "Primary System Privileges", "Security & Access Scope"),
        ("Women / Primary User", "Stealth SOS dispatch, Cab & Walk Guardian, Cybercrime reporting, Personal Vault access.", "Full read/write of own telemetry, contacts, and personal cryptographic evidence."),
        ("Trusted Guardians", "Receive real-time waypoint streams, route deviation push alerts, automated SMS links.", "Scoped read-only access to linked user's active SOS sessions and live transit breadcrumbs."),
        ("Police / Emergency Responders", "View live emergency case map, prioritize incidents, inspect victim profile & court dossier.", "Authorized access to active emergency distress sessions within geographic jurisdiction."),
        ("System Administrators", "Manage RBAC roles, audit log viewer, unsafe zone polygons, AI threat model parameters.", "Administrative system control, platform compliance, security monitoring, and server health.")
    ]
    for r_idx, row in enumerate(rb_data):
        for c_idx, val in enumerate(row):
            rbt.cell(r_idx, c_idx).text = val
    format_table(rbt, [Inches(2.5), Inches(5.2), Inches(4.0)], font_size=Pt(9.5))

    # ==========================================
    # SLIDE 5: Core Functional Requirements (BRD / SRS)
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "CORE FUNCTIONAL REQUIREMENTS (BRD / SRS)", "Engineered Capabilities for Life-Safety Operations")

    fr_table_shape = s5.shapes.add_table(6, 4, Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.7))
    frt = fr_table_shape.table
    fr_data = [
        ("Requirement ID", "Functional Feature", "Engineering Specification", "Verification Criteria"),
        ("FR-01", "Instant SOS Dispatch", "Emergency signal processing from trigger to server acknowledgment in <= 2 seconds.", "Sub-2s benchmark under 4G/LTE & BLE relay triggers."),
        ("FR-02", "Live GPS Waypoint Streams", "Continuous waypoint capture at 3–5s intervals with speed, heading, and Kalman smoothing.", "Live breadcrumbs rendered on guardian/responder maps without lag."),
        ("FR-03", "Route Deviation Alerts", "Heuristic Haversine algorithm flagging transit offsets >30m from planned route vector.", "Automated alert push to guardians when cab veers off polyline path."),
        ("FR-04", "Cyber Incident Logging", "Tesseract/EasyOCR parsing of screenshots and NLP threat severity score calculation.", "Accurate text extraction and risk score assignment in <3 seconds."),
        ("FR-05", "Automated Evidence Hashing", "Client-side SHA-256 fingerprinting stamped on all video/audio assets prior to upload.", "Cryptographic hash match verification between client and vault record.")
    ]
    for r_idx, row in enumerate(fr_data):
        for c_idx, val in enumerate(row):
            frt.cell(r_idx, c_idx).text = val
    format_table(frt, [Inches(1.5), Inches(2.8), Inches(4.6), Inches(2.8)], font_size=Pt(10))

    # ==========================================
    # SLIDE 6: Non-Functional Requirements & Security Mandates
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "NON-FUNCTIONAL REQUIREMENTS & SECURITY MANDATES", "Reliability, Privacy, Fault-Tolerance & Integrity Standards")

    nfr_cards = [
        ("NFR-01: High Availability (99.9% Uptime SLA)",
         "• Critical emergency endpoints hosted on auto-scaling WSGI workers with Redis queueing.\n• Zero single-point-of-failure: Health check monitoring with automatic pod failovers.\n• Rapid recovery objective (RTO < 30s) ensuring continuous 24/7 emergency dispatch availability.",
         Inches(0.8), Inches(1.3), Inches(5.7), Inches(2.4), (245, 240, 255), (147, 112, 219), (76, 29, 149)),
        ("NFR-02: Location Privacy & Data Minimization",
         "• Strict adherence to DPDP Act (India) & GDPR guidelines governing sensitive personal telemetry.\n• Automated 30-day scheduled purging of regular transit breadcrumbs for non-emergency sessions.\n• Geolocation coordinates encrypted in transit (TLS 1.3) and restricted strictly to linked guardians.",
         Inches(6.8), Inches(1.3), Inches(5.7), Inches(2.4), (238, 242, 255), (129, 140, 248), (49, 46, 129)),
        ("NFR-03: Offline SMS Fallback Queue",
         "• Automatic detection of cellular data packet loss or server unreachability on client devices.\n• Local SQLite coordinate queueing paired with automated SMS fallback gateway dispatch.\n• Direct SMS contains timestamp, device battery health, and live Google Maps coordinate hyperlink.",
         Inches(0.8), Inches(3.9), Inches(5.7), Inches(2.4), (255, 247, 237), (251, 146, 60), (194, 65, 12)),
        ("NFR-04: SHA-256 Evidence Integrity",
         "• Tamper-proof digital fingerprinting guaranteeing legal admissibility under Section 65B of Evidence Act.\n• Media digests calculated locally prior to transmission; any bit modification invalidates integrity.\n• Immutable database audit trail logging asset hash, file size, timestamp, and uploading UID.",
         Inches(6.8), Inches(3.9), Inches(5.7), Inches(2.4), (253, 244, 255), (232, 121, 249), (112, 26, 117))
    ]
    for title, content, x, y, w, h, bg, border, text_c in nfr_cards:
        box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(*bg)
        box.line.color.rgb = RGBColor(*border)
        box.line.width = Pt(1.5)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.15)
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Arial"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(*text_c)
        p2 = tf.add_paragraph()
        p2.text = content
        p2.font.name = "Arial"
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = C_DARK

    out_s6 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.45), Inches(11.7), Inches(0.65))
    out_s6.fill.solid()
    out_s6.fill.fore_color.rgb = RGBColor(241, 245, 249)
    out_s6.line.color.rgb = RGBColor(203, 213, 225)
    tf = out_s6.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Security Mandate : Security by design enforces zero-trust token authentication, ephemeral location retention, and immutable evidence notarization."
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(30, 41, 59)

    # ==========================================
    # SLIDE 7: High-Level System Architecture (4-Tier Design)
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "HIGH-LEVEL SYSTEM ARCHITECTURE (4-TIER DESIGN)", "Decoupled Service-Oriented Framework for Scalable Emergency Operations")

    arch_img = "extracted_assets/she_shield_4tier_architecture.png"
    if os.path.exists(arch_img):
        s7.shapes.add_picture(arch_img, Inches(0.8), Inches(1.3), width=Inches(7.2))

    # Architecture explanation (Right)
    tb_arch = s7.shapes.add_textbox(Inches(8.2), Inches(1.3), Inches(4.3), Inches(5.8))
    tf = tb_arch.text_frame
    tf.word_wrap = True
    arch_points = [
        ("Tier 1: Presentation Layer", "Flutter cross-platform client for Android/iOS with 60fps UI, stealth triggers, and sensors, coupled with React.js police command desk."),
        ("Tier 2: API Gateway & Security", "JWT authentication filter, RBAC role-enforcement, IP rate limiting, TLS 1.3 encryption, and reverse-proxy routing."),
        ("Tier 3: Core Application & AI/OCR", "Python Flask microservices handling parallel SOS dispatch, Haversine route deviation, PyTorch NLP threat classifier, and OCR screenshot pipeline."),
        ("Tier 4: Data & Storage Layer", "PostgreSQL for ACID relational records (users, RBAC, complaints, audit logs), Redis for sub-ms GPS telemetry streams, and AES-256 cloud media vault.")
    ]
    for i, (title, desc) in enumerate(arch_points):
        p1 = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Arial"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = C_PURPLE
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = "Arial"
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_DARK
        p_sp = tf.add_paragraph()
        p_sp.text = ""
        p_sp.font.size = Pt(4)

    # ==========================================
    # SLIDE 8: Architectural Diagrams & Data Pipelines
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "ARCHITECTURAL DIAGRAMS & DATA PIPELINES", "Emergency Parallel Broadcast & Cryptographic Evidence Custody Chain")

    pipe_img = "extracted_assets/she_shield_data_pipelines.png"
    if os.path.exists(pipe_img):
        s8.shapes.add_picture(pipe_img, Inches(0.8), Inches(1.3), width=Inches(11.7))

    # ==========================================
    # SLIDE 9: Technology Stack Justification
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "TECHNOLOGY STACK JUSTIFICATION", "Engineering Rationale for Architectural Choices")

    tech_table_shape = s9.shapes.add_table(6, 4, Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.7))
    techt = tech_table_shape.table
    tech_data = [
        ("Technology Component", "Category", "Selected Technology", "Engineering Justification in SheShield"),
        ("Mobile Client", "Presentation", "Flutter (Dart 3.x)", "Native 60fps performance on Android & iOS, unified single codebase, low-latency BLE sensor access."),
        ("Command Portal", "Presentation", "React.js + Tailwind", "High-density responsive web console for emergency police dispatchers with real-time WebSocket state."),
        ("Backend Services", "Application APIs", "Python Flask Microservices", "Lightweight WSGI services, native integration with PyTorch/OpenCV AI models, asynchronous worker queues."),
        ("Relational Storage", "Database", "PostgreSQL", "Strict ACID compliance, relational integrity for RBAC roles, audit trails, and structured police complaints."),
        ("In-Memory Telemetry", "Cache & Streams", "Redis Pub/Sub", "Sub-millisecond latency for live GPS waypoint broadcasting, temporary session caching, and rate limiting.")
    ]
    for r_idx, row in enumerate(tech_data):
        for c_idx, val in enumerate(row):
            techt.cell(r_idx, c_idx).text = val
    format_table(techt, [Inches(2.0), Inches(1.8), Inches(2.6), Inches(5.3)], font_size=Pt(10))

    # ==========================================
    # SLIDE 10: Security Architecture & Cryptographic Integrity
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "SECURITY ARCHITECTURE & CRYPTOGRAPHIC INTEGRITY", "Multi-Layered Cyber Defense and Chain-of-Custody Assurance")

    sec_cards = [
        ("1. Role-Based Access Control (RBAC) via JWT",
         "• Cryptographically signed JSON Web Tokens (JWT) enforcing granular endpoint permissions.\n• Strict role boundaries: Primary Users cannot modify admin zones; Guardians cannot view unlinked telemetry.\n• Short-lived access tokens (15 mins) paired with secure refresh tokens to prevent credential hijacking.",
         Inches(0.8), Inches(1.3), Inches(5.7), Inches(2.4), (245, 240, 255), (147, 112, 219), (76, 29, 149)),
        ("2. AES-256-GCM Vault Encryption",
         "• Authenticated AES-256 Galois/Counter Mode (GCM) encrypting all vaulted audio, video, and screenshot assets.\n• Unique cryptographic initialization vectors (IVs) generated per asset to prevent replay and leakage attacks.\n• Decryption keys restricted strictly to authorized forensic responders and subpoena-validated officials.",
         Inches(6.8), Inches(1.3), Inches(5.7), Inches(2.4), (238, 242, 255), (129, 140, 248), (49, 46, 129)),
        ("3. IP-Level & Token-Level Rate Limiting",
         "• Sliding-window rate limiters protecting emergency dispatch endpoints against Distributed Denial of Service (DDoS).\n• Redis-backed request tracking throttling unauthorized endpoint flooding while prioritizing active SOS signals.\n• Enforces IP reputation filtering and brute-force mitigation on authentication endpoints.",
         Inches(0.8), Inches(3.9), Inches(5.7), Inches(2.4), (255, 247, 237), (251, 146, 60), (194, 65, 12)),
        ("4. Cryptographic Hash Admissibility (Sec 65B)",
         "• SHA-256 digital fingerprint calculated client-side before media bytes leave user hardware.\n• Hash recorded in immutable ledger alongside UTC timestamp and device ID; bit-for-bit court proof.\n• Eliminates forensic defense challenges regarding evidence tampering in criminal proceedings.",
         Inches(6.8), Inches(3.9), Inches(5.7), Inches(2.4), (253, 244, 255), (232, 121, 249), (112, 26, 117))
    ]
    for title, content, x, y, w, h, bg, border, text_c in sec_cards:
        box = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(*bg)
        box.line.color.rgb = RGBColor(*border)
        box.line.width = Pt(1.5)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.15)
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Arial"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(*text_c)
        p2 = tf.add_paragraph()
        p2.text = content
        p2.font.name = "Arial"
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = C_DARK

    out_s10 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.45), Inches(11.7), Inches(0.65))
    out_s10.fill.solid()
    out_s10.fill.fore_color.rgb = RGBColor(240, 253, 250)
    out_s10.line.color.rgb = RGBColor(45, 212, 191)
    tf = out_s10.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Cryptographic Assurance : SheShield delivers military-grade confidentiality, tamper-evident forensic validity, and zero-trust authorization."
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 118, 110)

    # ==========================================
    # SLIDE 11: Current Implementation Progress (Weeks 1–11)
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "CURRENT IMPLEMENTATION PROGRESS (WEEKS 1–11)", "Sprint Deliverables and GitHub Repository Status")

    prog_table_shape = s11.shapes.add_table(7, 4, Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.7))
    progt = prog_table_shape.table
    prog_data = [
        ("Sprint Milestone", "Delivered Features & Modules", "Repository Artifacts (GitHub)", "Current Status"),
        ("Milestone 1: Foundations", "System architecture, UI wireframes, theme setup, Provider state framework.", "lib/main.dart, lib/theme/app_theme.dart", "100% Completed"),
        ("Milestone 2: Auth & RBAC", "User, Guardian, Responder, Admin registration and role-based route guard.", "lib/screens/auth_screens.dart, police_admin_screens.dart", "100% Completed"),
        ("Milestone 3: SOS Center", "Multi-modal trigger engine, 5s cancellable timer, volume button trigger, BLE sync.", "lib/screens/sos_screens.dart, lib/services/sos_service.dart", "100% Completed"),
        ("Milestone 4: Travel Monitor", "Cab Guardian & Walk Guardian with 30m route offset deviation calculation.", "lib/screens/travel_monitor_screen.dart, database_service.dart", "100% Completed"),
        ("Milestone 5: Cybercrime Desk", "Cyber complaint intake, OCR screenshot scan, PyTorch threat score, PDF export.", "lib/screens/cyber_screens.dart, pdf_report_service.dart", "100% Completed"),
        ("Milestone 6: Command Hub", "Admin Command Center, Responder live radar, and Evidence Locker vault.", "lib/screens/admin_command_center_screen.dart, vault_screens.dart", "100% Completed")
    ]
    for r_idx, row in enumerate(prog_data):
        for c_idx, val in enumerate(row):
            progt.cell(r_idx, c_idx).text = val
    format_table(progt, [Inches(2.2), Inches(4.5), Inches(3.2), Inches(1.8)], font_size=Pt(10))

    # ==========================================
    # SLIDE 12: Next Steps (Weeks 12–16) & Q&A Conclusion
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "NEXT STEPS (WEEKS 12–16) & CONCLUSION", "Project Roadmap, Engineering Milestones & Q&A Floor")

    roadmap_cards = [
        ("Weeks 12–13: End-to-End Integration",
         "• Full integration testing between Flutter mobile client, Flask microservices, and React admin desk.\n• WebSocket live synchronization hardening across simulated multi-node responder environments.\n• Edge-case verification for offline network reconnects and SQLite synchronization queues."),
        ("Week 14: GPS Geofencing Benchmarks",
         "• Field trials evaluating Haversine route deviation accuracy across dense urban corridors.\n• Kalman filter calibration to eliminate multipath GPS jitter around high-rise infrastructure.\n• Latency benchmarking for automated deviation alerts dispatched to guardian devices."),
        ("Week 15: High-Concurrency Load Testing",
         "• Locust and JMeter load testing simulating 5,000+ simultaneous emergency SOS trigger bursts.\n• Security penetration testing, API vulnerability scanning, and token lifecycle audits.\n• Performance optimization of PostgreSQL query indices and Redis pub/sub memory allocation."),
        ("Week 16: Release & Viva Evaluation",
         "• Production Docker containerization, staging deployment, and signed APK release.\n• Community pilot testing with feedback collection and final documentation synthesis.\n• Project Viva Voce defense for IT644 (WSSOA) Autumn 2026 evaluation.")
    ]
    y_pos = 1.3
    for title, content in roadmap_cards:
        box = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y_pos), Inches(11.7), Inches(1.05))
        box.fill.solid()
        box.fill.fore_color.rgb = C_PURPLE_LIGHT
        box.line.color.rgb = C_BORDER
        box.line.width = Pt(1.5)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.1)
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Arial"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = C_PURPLE
        p2 = tf.add_paragraph()
        p2.text = content
        p2.font.name = "Arial"
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = C_DARK
        y_pos += 1.15

    # Bottom Thank You & Q&A banner
    ty_box = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.0), Inches(11.7), Inches(1.1))
    ty_box.fill.solid()
    ty_box.fill.fore_color.rgb = RGBColor(76, 29, 149)
    tf = ty_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "THANK YOU  ·  QUESTIONS & ANSWERS"
    p.font.name = "Arial"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = "SHE SHIELD — Cyber-Integrated Safety Platform for Women  |  Dhirubhai Ambani University, Gandhinagar"
    p2.font.name = "Arial"
    p2.font.size = Pt(12)
    p2.font.color.rgb = RGBColor(233, 213, 255)
    p2.alignment = PP_ALIGN.CENTER

    output_path = "SheShield_CyberIntegrated_SafetyPlatform.pptx"
    prs.save(output_path)
    print(f"Successfully generated {output_path} with {len(prs.slides)} slides!")

if __name__ == "__main__":
    build_presentation()
